"""Phase 4.2: the four demo workbooks, every number from the truth ledger (ADR-013), every
derived figure as an Excel formula (cached values come from `recalc.sh`, LibreOffice
headless — openpyxl writes formulas only). Layout constants here are the single source the
validator (`validate_excel.py`) and the backend's predefined functions rely on: named ranges
are the contract, cell addresses are an implementation detail behind them.

    python -m seed_data.generator.generate_excel   # -> seed_data/excel/*.xlsx + manifest.json
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from seed_data.generator import debt_math
from seed_data.generator import facts as facts_mod

EXCEL_DIR = Path(__file__).resolve().parents[1] / "excel"
META_SHEET = "_meta"
BANNER = "DEMO — synthetic demo workbook (Company AI V0); figures are fictional"

FINANCIAL_MODEL = "Financial_Model_2026.xlsx"
COVENANT_REPORT = "Covenant_Report.xlsx"
BUDGET_VS_ACTUAL = "Budget_vs_Actual_2026.xlsx"
MONTHLY_PRODUCTION = "Monthly_Production_2026.xlsx"

# Fixed cell positions in each quarter sheet of the Covenant Report (SPEC_04 §5's example
# `Covenant_Report.xlsx Q2_2026!D14` is literally the DSCR cell).
Q_SHEET_CELLS = {
    "period": "D3",
    "cfads": "D5",
    "principal_half": "D6",
    "interest_half": "D7",
    "debt_service_quarter": "D8",
    "covenant": "D10",
    "outstanding": "D12",
    "dscr": "D14",
    "result": "D15",
}

_BOLD = Font(bold=True)


@dataclass(frozen=True)
class WorkbookSpec:
    file: str
    doc_id: str
    title: str
    document_type: str
    department: str
    subdepartment: str | None
    language: str
    document_date: str
    version_label: str
    sheets: tuple[str, ...]


def _ankara(raws: dict[str, Any]) -> dict[str, Any]:
    return raws["ankara_res"]["project"]


def _schedule(raws: dict[str, Any]) -> list[debt_math.ScheduleRow]:
    fin = _ankara(raws)["finance"]
    return debt_math.build_schedule(
        drawdowns=[(d["date"], float(d["amount"]["value"])) for d in fin["drawdowns"]],
        instalments=[
            (i["date"], float(i["principal"]["value"])) for i in fin["repayment_schedule"]
        ],
        base_rate_pct_by_year={
            b["year"]: float(b["rate_pct"]) for b in fin["base_rate_pct_by_year"]
        },
        margin_pct=float(fin["interest"]["margin_pct"]["value"]),
    )


def _quarter_end(period: str) -> date:
    quarter, year = period.split("_")
    month = {"Q1": 3, "Q2": 6, "Q3": 9, "Q4": 12}[quarter]
    day = 31 if month in (3, 12) else 30
    return date(int(year), month, day)


def _name(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    ref = f"'{sheet}'!${cell[0]}${cell[1:]}" if cell[1:].isdigit() else f"'{sheet}'!{cell}"
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _header(ws: Worksheet, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = _BOLD


def _meta(wb: Workbook, raws: dict[str, Any]) -> None:
    ws = wb.create_sheet(META_SHEET)
    ws["A1"] = BANNER
    ws["A2"] = "generated_at"
    ws["B2"] = datetime.now(UTC).replace(microsecond=0).isoformat()
    ws["A3"] = "ledger_schema_version"
    ws["B3"] = raws["ankara_res"]["meta"]["schema_version"]
    ws["A4"] = "demo_today"
    ws["B4"] = raws["ankara_res"]["meta"]["demo_today"].isoformat()
    ws.sheet_state = "hidden"
    # ADR-026: the day this workbook's snapshot values (e.g. `Outstanding_DemoToday`) were
    # computed for — read back at query time so a stale snapshot is refused, not served.
    _name(wb, "Ledger_DemoToday", META_SHEET, "B4")


# ---------------------------------------------------------------- 1. Financial Model


def build_financial_model(raws: dict[str, Any]) -> Workbook:
    project = _ankara(raws)
    fin = project["finance"]
    meta = raws["ankara_res"]["meta"]
    schedule = _schedule(raws)
    wb = Workbook()

    # Inputs
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"], ws["B1"] = "Project", project["name"]
    ws["A2"], ws["B2"] = "Capex (EUR)", fin["capex"]["value"]
    ws["A3"], ws["B3"] = "Equity (EUR)", fin["equity"]["value"]
    ws["A4"], ws["B4"] = "Total debt (EUR)", fin["total_debt"]["value"]
    ws["A5"], ws["B5"] = "Local bank debt (EUR)", fin["local_debt"]["value"]
    ws["A6"], ws["B6"] = "ECA debt (EUR)", fin["eca_debt"]["value"]
    ws["A7"], ws["B7"] = "Margin over base (%)", fin["interest"]["margin_pct"]["value"]
    ws["A8"], ws["B8"] = "Grace (months)", fin["grace_months"]["value"]
    ws["A9"], ws["B9"] = "Tenor initial (years)", fin["tenor_years"]["initial"]["value"]
    ws["A10"], ws["B10"] = "Tenor current (years)", fin["tenor_years"]["current"]["value"]
    ws["A11"], ws["B11"] = "Financial close", project["timeline"]["financial_close"]["date"]
    ws["A12"], ws["B12"] = "DEMO_TODAY", meta["demo_today"]
    ws["A13"], ws["B13"] = "Covenant DSCR initial", fin["dscr_covenant"]["initial"]["value"]
    ws["A14"], ws["B14"] = "Covenant DSCR current", fin["dscr_covenant"]["current"]["value"]
    ws["A15"], ws["B15"] = "Covenant change effective", date(2025, 3, 15)
    for c in ("B11", "B12", "B15"):
        ws[c].number_format = "DD.MM.YYYY"
    _header(ws, 17, ["Year", "Base rate % (EURIBOR 6M, fictional)"])
    base_first = 18
    rates = sorted(fin["base_rate_pct_by_year"], key=lambda b: b["year"])
    for i, b in enumerate(rates):
        ws.cell(row=base_first + i, column=1, value=b["year"])
        ws.cell(row=base_first + i, column=2, value=float(b["rate_pct"]))
    base_last = base_first + len(rates) - 1
    draw_first = base_last + 3
    _header(ws, draw_first - 1, ["Drawdown date", "Amount (EUR)"])
    for i, d in enumerate(fin["drawdowns"]):
        ws.cell(row=draw_first + i, column=1, value=d["date"]).number_format = "DD.MM.YYYY"
        ws.cell(row=draw_first + i, column=2, value=d["amount"]["value"])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    _name(wb, "Inputs_TotalDebt", "Inputs", "B4")
    _name(wb, "Inputs_Margin", "Inputs", "B7")
    _name(wb, "Inputs_DemoToday", "Inputs", "B12")
    base_table = f"Inputs!$A${base_first}:$B${base_last}"

    # Debt schedule — formulas chained row to row, rates looked up from Inputs.
    ws = wb.create_sheet("Debt")
    _header(
        ws,
        1,
        [
            "Period",
            "End date",
            "Opening (EUR)",
            "Drawdown (EUR)",
            "Principal (EUR)",
            "All-in rate %",
            "Interest (EUR)",
            "Closing (EUR)",
        ],
    )
    demo_today = meta["demo_today"]
    outstanding_row: int | None = None
    row_by_period: dict[str, int] = {}
    for i, r in enumerate(schedule):
        row = 2 + i
        row_by_period[r.period] = row
        ws.cell(row=row, column=1, value=r.period)
        ws.cell(row=row, column=2, value=r.end).number_format = "DD.MM.YYYY"
        ws.cell(row=row, column=3, value=0 if i == 0 else f"=H{row - 1}")
        ws.cell(row=row, column=4, value=r.drawdown)
        ws.cell(row=row, column=5, value=r.principal)
        ws.cell(
            row=row, column=6, value=f"=VLOOKUP(YEAR(B{row}),{base_table},2,FALSE)+Inputs_Margin"
        )
        ws.cell(row=row, column=7, value=f"=ROUND(C{row}*F{row}/100/2,2)")
        ws.cell(row=row, column=8, value=f"=C{row}+D{row}-E{row}")
        if r.end <= demo_today:
            outstanding_row = row
    assert outstanding_row is not None
    _name(wb, "Outstanding_DemoToday", "Debt", f"H{outstanding_row}")
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 16
    debt_last = 1 + len(schedule)

    # DSCR — quarterly: cfads (ledger) / half-year service ÷ 2 (formula into Debt).
    ws = wb.create_sheet("DSCR")
    _header(
        ws,
        1,
        [
            "Quarter",
            "Quarter end",
            "Half",
            "CFADS (EUR)",
            "Debt service (EUR)",
            "DSCR",
            "Covenant",
            "Result",
        ],
    )
    cfads = {c["period"]: c["cfads"]["value"] for c in fin["cfads_by_quarter"]}
    for i, test in enumerate(fin["covenant_tests"]):
        row = 2 + i
        period = test["period"]
        half = debt_math.half_of_quarter(period)
        ws.cell(row=row, column=1, value=period)
        ws.cell(row=row, column=2, value=_quarter_end(period)).number_format = "DD.MM.YYYY"
        ws.cell(row=row, column=3, value=half)
        ws.cell(row=row, column=4, value=cfads[period])
        ws.cell(
            row=row,
            column=5,
            value=(
                f"=(INDEX(Debt!$E$2:$E${debt_last},MATCH(C{row},Debt!$A$2:$A${debt_last},0))"
                f"+INDEX(Debt!$G$2:$G${debt_last},MATCH(C{row},Debt!$A$2:$A${debt_last},0)))/2"
            ),
        )
        ws.cell(row=row, column=6, value=f"=ROUND(D{row}/E{row},2)")
        ws.cell(row=row, column=7, value=f"=IF(B{row}>=Inputs!$B$15,Inputs!$B$14,Inputs!$B$13)")
        ws.cell(row=row, column=8, value=f'=IF(F{row}>=G{row},"pass","fail")')
        _name(wb, f"DSCR_{period}", "DSCR", f"F{row}")
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 16

    # Cashflow
    ws = wb.create_sheet("Cashflow")
    _header(
        ws,
        1,
        [
            "Quarter",
            "CFADS (EUR)",
            "Debt service (EUR)",
            "Cash to equity (EUR)",
            "Cumulative (EUR)",
        ],
    )
    for i, test in enumerate(fin["covenant_tests"]):
        row = 2 + i
        ws.cell(row=row, column=1, value=test["period"])
        ws.cell(row=row, column=2, value=f"=DSCR!D{row}")
        ws.cell(row=row, column=3, value=f"=DSCR!E{row}")
        ws.cell(row=row, column=4, value=f"=B{row}-C{row}")
        ws.cell(row=row, column=5, value=f"=D{row}" if i == 0 else f"=E{row - 1}+D{row}")
    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 20
    _meta(wb, raws)
    return wb


# ---------------------------------------------------------------- 2. Covenant Report


def build_covenant_report(raws: dict[str, Any]) -> Workbook:
    fin = _ankara(raws)["finance"]
    schedule = _schedule(raws)
    by_half = {r.period: r for r in schedule}
    cfads = {c["period"]: c["cfads"]["value"] for c in fin["cfads_by_quarter"]}
    change_effective = date(2025, 3, 15)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _header(summary, 1, ["Quarter", "DSCR", "Covenant", "Result", "Outstanding after half (EUR)"])

    for i, test in enumerate(fin["covenant_tests"]):
        period = test["period"]
        half = by_half[debt_math.half_of_quarter(period)]
        covenant = (
            fin["dscr_covenant"]["current"]["value"]
            if _quarter_end(period) >= change_effective
            else fin["dscr_covenant"]["initial"]["value"]
        )
        ws = wb.create_sheet(period)
        ws["A1"] = f"Covenant compliance test — {period}"
        ws["A1"].font = _BOLD
        labels = {
            "period": ("Period", period),
            "cfads": ("CFADS (EUR)", cfads[period]),
            "principal_half": ("Principal, half-year (EUR)", half.principal),
            "interest_half": ("Interest, half-year (EUR)", half.interest),
            "debt_service_quarter": (
                "Debt service, quarter (EUR) = (principal + interest) / 2",
                "=(D6+D7)/2",
            ),
            "covenant": ("Covenant threshold (min DSCR)", covenant),
            "outstanding": ("Outstanding after half (EUR)", half.closing),
            "dscr": ("DSCR = CFADS / debt service", "=ROUND(D5/D8,2)"),
            "result": ("Result", '=IF(D14>=D10,"pass","fail")'),
        }
        for key, (label, value) in labels.items():
            cell = Q_SHEET_CELLS[key]
            ws[f"B{cell[1:]}"] = label
            ws[cell] = value
        ws.column_dimensions["B"].width = 52
        ws.column_dimensions["D"].width = 18
        _name(wb, f"DSCR_{period}", period, Q_SHEET_CELLS["dscr"])
        _name(wb, f"Outstanding_{period}", period, Q_SHEET_CELLS["outstanding"])

        row = 2 + i
        summary.cell(row=row, column=1, value=period)
        summary.cell(row=row, column=2, value=f"='{period}'!{Q_SHEET_CELLS['dscr']}")
        summary.cell(row=row, column=3, value=f"='{period}'!{Q_SHEET_CELLS['covenant']}")
        summary.cell(row=row, column=4, value=f"='{period}'!{Q_SHEET_CELLS['result']}")
        summary.cell(row=row, column=5, value=f"='{period}'!{Q_SHEET_CELLS['outstanding']}")
    for col in range(1, 6):
        summary.column_dimensions[get_column_letter(col)].width = 22
    _meta(wb, raws)
    return wb


# ---------------------------------------------------------------- 3. Budget vs Actual


def build_budget_vs_actual(raws: dict[str, Any]) -> Workbook:
    rows = _ankara(raws)["operations"]["budget_vs_actual"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _header(ws, 1, ["Quarter", "Budget (TRY)", "Actual (TRY)", "Variance (TRY)", "Variance %"])
    for i, r in enumerate(rows):
        row = 2 + i
        ws.cell(row=row, column=1, value=r["period"])
        ws.cell(row=row, column=2, value=r["budget"]["value"])
        ws.cell(row=row, column=3, value=r["actual"]["value"])
        ws.cell(row=row, column=4, value=f"=C{row}-B{row}")
        ws.cell(row=row, column=5, value=f"=ROUND(D{row}/B{row},4)")
        _name(wb, f"Budget_{r['period']}", "Summary", f"B{row}")
        _name(wb, f"Actual_{r['period']}", "Summary", f"C{row}")
        _name(wb, f"Variance_{r['period']}", "Summary", f"D{row}")
    total = 2 + len(rows)
    ws.cell(row=total, column=1, value="Total").font = _BOLD
    ws.cell(row=total, column=2, value=f"=SUM(B2:B{total - 1})")
    ws.cell(row=total, column=3, value=f"=SUM(C2:C{total - 1})")
    ws.cell(row=total, column=4, value=f"=C{total}-B{total}")
    ws.cell(row=total, column=5, value=f"=ROUND(D{total}/B{total},4)")
    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 18
    _meta(wb, raws)
    return wb


# ---------------------------------------------------------------- 4. Monthly Production


def build_monthly_production(raws: dict[str, Any]) -> Workbook:
    project = _ankara(raws)
    rows = project["operations"]["monthly_production"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Production"
    _header(ws, 1, ["Month", "MWh", "Availability %", "Capacity factor %"])
    for i, r in enumerate(rows):
        row = 2 + i
        ws.cell(row=row, column=1, value=r["month"])
        ws.cell(row=row, column=2, value=r["mwh"])
        ws.cell(row=row, column=3, value=r["availability_pct"])
        ws.cell(row=row, column=4, value=r["capacity_factor_pct"])
        _name(wb, f"Production_{r['month'].replace('-', '_')}", "Production", f"B{row}")
    last = 1 + len(rows)
    for col in range(1, 5):
        ws.column_dimensions[get_column_letter(col)].width = 18

    kpi = wb.create_sheet("KPI")
    _header(kpi, 1, ["Year", "Total MWh", "Avg availability %", "Avg capacity factor %", "Months"])
    years = sorted({r["month"][:4] for r in rows})
    for i, year in enumerate(years):
        row = 2 + i
        kpi.cell(row=row, column=1, value=int(year))
        kpi.cell(
            row=row,
            column=2,
            value=f'=SUMIF(Production!$A$2:$A${last},"{year}-*",Production!$B$2:$B${last})',
        )
        kpi.cell(
            row=row,
            column=3,
            value=f'=ROUND(AVERAGEIF(Production!$A$2:$A${last},"{year}-*",Production!$C$2:$C${last}),1)',
        )
        kpi.cell(
            row=row,
            column=4,
            value=f'=ROUND(AVERAGEIF(Production!$A$2:$A${last},"{year}-*",Production!$D$2:$D${last}),1)',
        )
        kpi.cell(row=row, column=5, value=f'=COUNTIF(Production!$A$2:$A${last},"{year}-*")')
        _name(wb, f"MWh_Total_{year}", "KPI", f"B{row}")
    for col in range(1, 6):
        kpi.column_dimensions[get_column_letter(col)].width = 20
    kpi["A8"] = "Capacity (MW)"
    kpi["B8"] = project["capacity_mw"]["current"]["value"]
    _name(wb, "Capacity_MW", "KPI", "B8")
    _meta(wb, raws)
    return wb


# ---------------------------------------------------------------- manifest + entrypoint

SPECS: tuple[WorkbookSpec, ...] = (
    WorkbookSpec(
        FINANCIAL_MODEL,
        "DOC-ANK-FIN-008",
        "Financial Model 2026",
        "Financial Model",
        "finans",
        None,
        "en",
        "2026-09-01",
        "V2026",
        ("Inputs", "Debt", "DSCR", "Cashflow", META_SHEET),
    ),
    WorkbookSpec(
        COVENANT_REPORT,
        "DOC-ANK-FIN-009",
        "Covenant Report (workbook)",
        "Covenant Report",
        "finans",
        None,
        "en",
        "2026-08-15",
        "Q2_2026",
        ("Summary", META_SHEET),
    ),
    WorkbookSpec(
        BUDGET_VS_ACTUAL,
        "DOC-ANK-OPS-002",
        "Budget vs Actual 2026",
        "Budget vs Actual",
        "enerji_grubu",
        "enerji_bakim",
        "tr",
        "2026-07-31",
        "Q2_2026",
        ("Summary", META_SHEET),
    ),
    WorkbookSpec(
        MONTHLY_PRODUCTION,
        "DOC-ANK-OPS-003",
        "Monthly Production 2026",
        "Monthly Production",
        "enerji_grubu",
        "enerji_bakim",
        "tr",
        "2026-09-05",
        "2026-08",
        ("Production", "KPI", META_SHEET),
    ),
)

BUILDERS = {
    FINANCIAL_MODEL: build_financial_model,
    COVENANT_REPORT: build_covenant_report,
    BUDGET_VS_ACTUAL: build_budget_vs_actual,
    MONTHLY_PRODUCTION: build_monthly_production,
}


def _ledger_parties(raws: dict[str, Any], doc_id: str) -> list[str]:
    for ledger in raws.values():
        for doc in (ledger.get("documents") or []) if isinstance(ledger, dict) else []:
            if doc.get("id") == doc_id:
                return [str(p) for p in (doc.get("parties") or [])]
    return []


def _ledger_document(raws: dict[str, Any], doc_id: str) -> dict[str, Any]:
    for ledger in raws.values():
        for doc in (ledger.get("documents") or []) if isinstance(ledger, dict) else []:
            if doc.get("id") == doc_id:
                return dict(doc)
    raise KeyError(f"{doc_id} is not in the ledger")


def _ledger_confidentiality(raws: dict[str, Any], doc_id: str) -> str:
    """Single source of truth for the gate fields (ADR-004): the ledger's `confidentiality`
    — Soru 15 (09.10.2026) flipped the financial model to `normal` there, not in a spec."""
    return str(_ledger_document(raws, doc_id)["confidentiality"])


def generate(out_dir: Path = EXCEL_DIR) -> list[dict[str, Any]]:
    raws = facts_mod.load_raws()
    out_dir.mkdir(parents=True, exist_ok=True)
    entries: list[dict[str, Any]] = []
    spv = str(raws["ankara_res"]["project"]["spv"]["name"]["value"])
    for spec in SPECS:
        wb = BUILDERS[spec.file](raws)
        wb.save(out_dir / spec.file)
        entries.append(
            {
                "external_ref": spec.doc_id,
                "file": spec.file,
                "title": spec.title,
                "document_type": spec.document_type,
                "department": spec.department,
                "subdepartment": spec.subdepartment,
                "project_code": "ANK_RES",
                "counterparty": spv,
                "document_date": spec.document_date,
                "effective_date": None,
                "status": "active",
                "version_label": spec.version_label,
                "version_number": 1,
                "supersedes_ref": None,
                "related_refs": [],
                "tags": [],
                # B-28b: parties from the ledger document entry (SPV + bank where recorded).
                "parties": _ledger_parties(raws, spec.doc_id),
                "language": spec.language,
                "confidentiality": _ledger_confidentiality(raws, spec.doc_id),
                "source_type": "xlsx",
                "sheets": list(spec.sheets),
                "named_ranges": sorted(wb.defined_names.keys()),
            }
        )
    manifest = {
        "generated_at": datetime.now(UTC).replace(microsecond=0).isoformat(),
        "ledger_demo_today": raws["ankara_res"]["meta"]["demo_today"].isoformat(),
        "workbooks": entries,
    }
    (out_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return entries


if __name__ == "__main__":
    for entry in generate():
        print(
            f"{entry['file']}: {len(entry['named_ranges'])} named ranges, sheets {entry['sheets']}"
        )
