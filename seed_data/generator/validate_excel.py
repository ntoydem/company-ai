"""Phase 4.2 workbook checks (W1-W5) over the committed `seed_data/excel/*.xlsx`:
W1 every formula cell has a cached value (LibreOffice recalc happened),
W2 the named ranges the backend's predefined functions rely on exist,
W3 `_meta` is hidden,
W4 cached values reproduce the ledger (outstanding on demo_today, every quarter's DSCR,
   budget/actual rows, monthly MWh and the yearly totals),
W5 file and sheet names match the spec.
Part of `make lint` (fast, no LibreOffice needed — it validates, never recalculates)."""

from __future__ import annotations

import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from seed_data.generator import debt_math
from seed_data.generator import facts as facts_mod
from seed_data.generator.generate_excel import (
    BUDGET_VS_ACTUAL,
    COVENANT_REPORT,
    EXCEL_DIR,
    FINANCIAL_MODEL,
    META_SHEET,
    MONTHLY_PRODUCTION,
    Q_SHEET_CELLS,
    SPECS,
)


@dataclass
class Report:
    errors: list[str] = field(default_factory=list)

    def error(self, message: str) -> None:
        self.errors.append(message)


def _named(wb: Any, name: str) -> tuple[str, str]:
    defined = wb.defined_names[name]
    sheet, ref = next(iter(defined.destinations))
    return sheet, ref.replace("$", "")


def _value(wb_values: Any, wb: Any, name: str) -> Any:
    sheet, ref = _named(wb, name)
    return wb_values[sheet][ref].value


def check_workbook(path: Path, report: Report) -> None:
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    uncached = 0
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
                if is_formula and values[ws.title][cell.coordinate].value is None:
                    uncached += 1
    if uncached:
        report.error(
            f"{path.name}: W1 {uncached} formula cell(s) without a cached value — run recalc.sh"
        )
    if META_SHEET not in formulas.sheetnames or formulas[META_SHEET].sheet_state != "hidden":
        report.error(f"{path.name}: W3 `{META_SHEET}` missing or not hidden")


def check_ledger_consistency(excel_dir: Path, report: Report) -> None:
    raws = facts_mod.load_raws()
    project = raws["ankara_res"]["project"]
    fin = project["finance"]
    demo_today: date = raws["ankara_res"]["meta"]["demo_today"]
    schedule = debt_math.build_schedule(
        drawdowns=[(d["date"], float(d["amount"]["value"])) for d in fin["drawdowns"]],
        instalments=[
            (i["date"], float(i["principal"]["value"])) for i in fin["repayment_schedule"]
        ],
        base_rate_pct_by_year={
            b["year"]: float(b["rate_pct"]) for b in fin["base_rate_pct_by_year"]
        },
        margin_pct=float(fin["interest"]["margin_pct"]["value"]),
    )

    fm = load_workbook(excel_dir / FINANCIAL_MODEL, data_only=False)
    fm_values = load_workbook(excel_dir / FINANCIAL_MODEL, data_only=True)
    outstanding = _value(fm_values, fm, "Outstanding_DemoToday")
    expected = fin["outstanding_debt_as_of_demo_today"]["value"]
    if outstanding is None or abs(float(outstanding) - float(expected)) > 0.5:
        report.error(
            f"{FINANCIAL_MODEL}: W4 Outstanding_DemoToday {outstanding} != ledger {expected}"
        )
    if abs(debt_math.outstanding_on(schedule, demo_today) - float(expected)) > 0.5:
        report.error(f"{FINANCIAL_MODEL}: W4 debt_math outstanding differs from ledger")

    cr = load_workbook(excel_dir / COVENANT_REPORT, data_only=False)
    cr_values = load_workbook(excel_dir / COVENANT_REPORT, data_only=True)
    for test in fin["covenant_tests"]:
        period = test["period"]
        for name, wb, wb_values in (
            (FINANCIAL_MODEL, fm, fm_values),
            (COVENANT_REPORT, cr, cr_values),
        ):
            try:
                got = _value(wb_values, wb, f"DSCR_{period}")
            except KeyError:
                report.error(f"{name}: W2 named range DSCR_{period} missing")
                continue
            if got is None or abs(float(got) - float(test["dscr"])) > 0.005:
                report.error(f"{name}: W4 DSCR_{period} = {got}, ledger {test['dscr']}")
        cell = cr_values[period][Q_SHEET_CELLS["result"]].value
        if cell != test["result"]:
            report.error(
                f"{COVENANT_REPORT}: W4 {period}!{Q_SHEET_CELLS['result']} = {cell!r}, "
                f"ledger {test['result']!r}"
            )

    bva = load_workbook(excel_dir / BUDGET_VS_ACTUAL, data_only=False)
    bva_values = load_workbook(excel_dir / BUDGET_VS_ACTUAL, data_only=True)
    for row in project["operations"]["budget_vs_actual"]:
        period = row["period"]
        for kind in ("Budget", "Actual"):
            got = _value(bva_values, bva, f"{kind}_{period}")
            expected = row[kind.lower()]["value"]
            if got != expected:
                report.error(f"{BUDGET_VS_ACTUAL}: W4 {kind}_{period} = {got}, ledger {expected}")
        variance = _value(bva_values, bva, f"Variance_{period}")
        if variance != row["actual"]["value"] - row["budget"]["value"]:
            report.error(f"{BUDGET_VS_ACTUAL}: W4 Variance_{period} = {variance}")

    mp = load_workbook(excel_dir / MONTHLY_PRODUCTION, data_only=False)
    mp_values = load_workbook(excel_dir / MONTHLY_PRODUCTION, data_only=True)
    totals: dict[str, float] = {}
    for row in project["operations"]["monthly_production"]:
        got = _value(mp_values, mp, f"Production_{row['month'].replace('-', '_')}")
        if got != row["mwh"]:
            report.error(
                f"{MONTHLY_PRODUCTION}: W4 Production_{row['month']} = {got}, ledger {row['mwh']}"
            )
        totals[row["month"][:4]] = totals.get(row["month"][:4], 0.0) + float(row["mwh"])
    for year, total in totals.items():
        got = _value(mp_values, mp, f"MWh_Total_{year}")
        if got is None or abs(float(got) - total) > 0.5:
            report.error(f"{MONTHLY_PRODUCTION}: W4 MWh_Total_{year} = {got}, ledger sum {total}")


def validate(excel_dir: Path = EXCEL_DIR) -> Report:
    report = Report()
    for spec in SPECS:
        path = excel_dir / spec.file
        if not path.exists():
            report.error(f"{spec.file}: W5 missing — run `make excel`")
            continue
        wb = load_workbook(path, data_only=False)
        if set(spec.sheets) != set(wb.sheetnames) and spec.file != COVENANT_REPORT:
            report.error(f"{spec.file}: W5 sheets {wb.sheetnames} != spec {list(spec.sheets)}")
        if spec.file == COVENANT_REPORT and not {"Summary", META_SHEET} <= set(wb.sheetnames):
            report.error(f"{spec.file}: W5 Summary/_meta missing")
        check_workbook(path, report)
    if not report.errors:
        check_ledger_consistency(excel_dir, report)
    return report


def main() -> int:
    report = validate()
    for error in report.errors:
        print(f"ERROR {error}")
    print(f"{len(report.errors)} error(s)")
    return 1 if report.errors else 0


if __name__ == "__main__":
    sys.exit(main())
