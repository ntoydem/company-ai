"""`CalculationEngine` (SPEC_04 §6, ADR-011) and its only V0 implementation,
`CachedValueEngine`: cell values are openpyxl's cached values (`data_only=True`), sheets
become DuckDB tables (via Polars) for read-only SELECTs, named ranges give predefined
functions their cells *and* their citations. No server-side recalculation: a workbook with
formula cells that carry no cached value raises `NeedsRecalculationError`, and the user is
told to recalculate and save it in Excel."""

from __future__ import annotations

import logging
import math
import re
import threading
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Protocol, cast
from uuid import UUID

import duckdb
import polars as pl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.excel.inspect import WorkbookInfo, inspect_file, read_csv_rows
from app.excel.sql_guard import DEFAULT_ROW_LIMIT, guard_select

log = logging.getLogger(__name__)

DEFAULT_TIMEOUT_S = 10.0
ROW_COLUMN = "_row"  # Excel row number of each table row — the citation range comes from it
# ADR-037: the header is the first of the top rows with ≥ 3 non-empty cells, all text; rows
# above it (title, snapshot date, loan total) go to a `<table>_info` key/value table.
# Threshold 2 would take Covenant quarter sheets' `Period | Q2_2026` row as a header.
HEADER_SCAN_ROWS = 10
HEADER_MIN_TEXT_CELLS = 3
INFO_SUFFIX = "_info"


class ExcelError(Exception):
    """Base: something the user should hear about in Turkish (the API maps it)."""


class NeedsRecalculationError(ExcelError):
    """Formula cells without cached values (file saved by a tool that does not compute)."""


class QueryTimeoutError(ExcelError):
    """DuckDB query interrupted after the configured timeout."""


class UnknownNamedRangeError(ExcelError):
    pass


@dataclass(frozen=True)
class SourceRange:
    """`Covenant_Report.xlsx Q2_2026!D14` — file + sheet + range (SPEC_04 §5)."""

    file: str
    sheet: str
    ref: str
    document_id: UUID | None = None

    @property
    def label(self) -> str:
        return f"{self.file} {self.sheet}!{self.ref}"


@dataclass
class LoadedWorkbook:
    path: Path
    info: WorkbookInfo
    document_id: UUID | None
    file_name: str | None = None  # citation name; storage is always original.<ext>
    tables: dict[str, pl.DataFrame] = field(default_factory=dict)  # table name -> rows
    sheet_by_table: dict[str, str] = field(default_factory=dict)
    values: Any = None  # openpyxl workbook, data_only=True (None for csv)

    @property
    def file(self) -> str:
        return self.file_name or self.path.name


@dataclass(frozen=True)
class QueryResult:
    columns: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]
    sources: tuple[SourceRange, ...]
    truncated: bool


@dataclass(frozen=True)
class CalcResult:
    value: Any
    unit: str
    source: SourceRange
    detail: str = ""


class CalculationEngine(Protocol):
    def load(
        self, path: Path, *, document_id: UUID | None = None, file_name: str | None = None
    ) -> LoadedWorkbook: ...
    def run_sql(self, workbooks: list[LoadedWorkbook], sql: str) -> QueryResult: ...
    def read_named(self, workbook: LoadedWorkbook, name: str) -> tuple[Any, SourceRange]: ...


_IDENT = re.compile(r"[^a-z0-9_]+")


def table_name(file: str, sheet: str) -> str:
    stem = Path(file).stem
    return (
        f"{_IDENT.sub('_', stem.lower()).strip('_')}__{_IDENT.sub('_', sheet.lower()).strip('_')}"
    )


def _unique_columns(header: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    names: list[str] = []
    for i, raw in enumerate(header):
        base = _IDENT.sub("_", str(raw).lower()).strip("_") if raw not in (None, "") else ""
        base = base or f"col_{get_column_letter(i + 1).lower()}"
        if base in seen:
            seen[base] += 1
            base = f"{base}_{seen[base]}"
        else:
            seen[base] = 0
        names.append(base)
    return names


def _frame(header: list[Any], rows: list[tuple[int, list[Any]]]) -> pl.DataFrame:
    columns = _unique_columns(header)
    data: dict[str, list[Any]] = {ROW_COLUMN: [r for r, _ in rows]}
    for j, name in enumerate(columns):
        data[name] = [(values[j] if j < len(values) else None) for _, values in rows]
    try:
        return pl.DataFrame(data, strict=False)
    except pl.exceptions.ComputeError:
        # İş Emri 2026-10-12-F Adım 1: bir sütunun ilk hücreleri tarih/sayı sanılıp
        # sonraki bir satırda metin (örn. düzensiz bir önsöz satırı — başlık/anlık
        # görüntü tarihi — veri gibi okunduğunda) görülünce polars `strict=False`'a
        # rağmen hâlâ `ComputeError` fırlatıyor (ham Excel hücre değerleri, LLM çıktısı
        # gibi, biçimine güvenilmeyen bir girdi — aynı "ham metni varsay, istisna
        # fırlatma" ilkesi burada da geçerli). Tüm hücreler metne çevrilip tekrar
        # denenir; hiçbir veri kaybolmaz, yalnız o sütunun tipi sayısal/tarih değil
        # metin olur (SQL sorguları üzerinde CAST gerekebilir, ama 500 yerine kullanılabilir
        # bir tablo döner).
        coerced: dict[str, list[Any]] = {
            key: values if key == ROW_COLUMN else [None if v is None else str(v) for v in values]
            for key, values in data.items()
        }
        return pl.DataFrame(coerced, strict=False)


def _filled(value: Any) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def header_row(rows: list[tuple[Any, ...]]) -> int:
    """1-based header row among the sheet's first rows; 1 (the pre-ADR-037 rule) when no
    row has ≥ `HEADER_MIN_TEXT_CELLS` filled cells that are all text."""
    for number, row in enumerate(rows[:HEADER_SCAN_ROWS], start=1):
        cells = [v for v in row if _filled(v)]
        if len(cells) >= HEADER_MIN_TEXT_CELLS and all(isinstance(v, str) for v in cells):
            return number
    return 1


Rows = list[tuple[int, tuple[Any, ...]]]  # (Excel row number, cell values)


def _typed(value: Any) -> bool:
    return isinstance(value, int | float | date) and not isinstance(value, bool)


def split_footer(body: Rows, width: int) -> tuple[Rows, Rows, bool]:
    """ADR-037 (Naci 17.10): on a sheet whose header is below row 1, the rows after the
    first fully blank row under the data are notes ("Kontrol", "Not") — unless one of them
    looks like data (≥ half of the `width` header columns filled with numbers/dates); then
    nothing is cut and the third value is True ("kesilmedi, belirsiz")."""
    started = False
    for i, (_, row) in enumerate(body):
        filled = any(_filled(v) for v in row)
        if filled:
            started = True
        elif started:
            notes = [(n, r) for n, r in body[i + 1 :] if any(_filled(v) for v in r)]
            if not notes:
                return body, [], False
            data_like = math.ceil(width / 2)
            if any(sum(_typed(v) for v in r) >= data_like for _, r in notes):
                return body, [], True
            return body[:i], notes, False
    return body, [], False


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime) and value.time() == time.min:
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _info_frame(rows: Rows) -> pl.DataFrame | None:
    """Prelude and cut note rows as `key | value` text (a mixed value column would be
    coerced — a date became a float): first filled cell, then the next one."""
    pairs: list[tuple[int, list[Any]]] = []
    for number, row in rows:
        cells = [v for v in row if _filled(v)]
        if cells:
            pairs.append((number, [_text(cells[0]), _text(cells[1]) if len(cells) > 1 else None]))
    return _frame(["key", "value"], pairs) if pairs else None


class CachedValueEngine:
    def __init__(
        self, *, timeout_s: float = DEFAULT_TIMEOUT_S, row_limit: int = DEFAULT_ROW_LIMIT
    ) -> None:
        self._timeout_s = timeout_s
        self._row_limit = row_limit

    # ---- loading

    def load(
        self, path: Path, *, document_id: UUID | None = None, file_name: str | None = None
    ) -> LoadedWorkbook:
        info = inspect_file(path)
        if info.needs_recalculation:
            raise NeedsRecalculationError(file_name or path.name)
        loaded = LoadedWorkbook(path=path, info=info, document_id=document_id, file_name=file_name)
        if info.kind == "csv":
            rows = read_csv_rows(path)
            if rows:
                body = [(i + 2, list(r)) for i, r in enumerate(rows[1:])]
                stem = Path(loaded.file).stem
                name = table_name(loaded.file, stem)
                loaded.tables[name] = _frame(list(rows[0]), body)
                loaded.sheet_by_table[name] = stem
            return loaded
        values = load_workbook(path, data_only=True)  # keep_vba=False: macros never loaded
        loaded.values = values
        for ws in values.worksheets:
            if ws.sheet_state != "visible" or ws.max_row < 2:
                continue
            top = list(ws.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True))
            first = header_row(top)
            header = list(top[first - 1])
            start = first + 1
            # Excel row numbers must survive skipped blank rows (they are the citation).
            below: Rows = list(enumerate(ws.iter_rows(min_row=start, values_only=True), start))
            notes: Rows = []
            if first > 1:
                below, notes, kept = split_footer(below, sum(_filled(v) for v in header))
                if kept:
                    log.info("excel footer not cut", extra={"file": loaded.file, "sheet": ws.title})
            sheet_rows = [(n, list(row)) for n, row in below if any(v is not None for v in row)]
            if not sheet_rows:
                continue
            name = table_name(loaded.file, ws.title)
            loaded.tables[name] = _frame(header, sheet_rows)
            loaded.sheet_by_table[name] = ws.title
            prelude = _info_frame(list(enumerate(top[: first - 1], start=1)) + notes)
            if prelude is not None:
                loaded.tables[name + INFO_SUFFIX] = prelude
                loaded.sheet_by_table[name + INFO_SUFFIX] = ws.title
        return loaded

    # ---- SQL

    def run_sql(self, workbooks: list[LoadedWorkbook], sql: str) -> QueryResult:
        tables: dict[str, tuple[LoadedWorkbook, pl.DataFrame]] = {}
        for wb in workbooks:
            for name, frame in wb.tables.items():
                tables[name] = (wb, frame)
        guarded = guard_select(sql, set(tables), row_limit=self._row_limit)

        conn = duckdb.connect(database=":memory:", config={"enable_external_access": False})
        try:
            for name in guarded.tables:
                conn.register(name, tables[name][1])
            timer = threading.Timer(self._timeout_s, conn.interrupt)
            timer.start()
            try:
                cursor = conn.execute(guarded.sql)
                columns = tuple(d[0] for d in cursor.description or ())
                rows = tuple(tuple(r) for r in cursor.fetchall())
            except duckdb.InterruptException as exc:
                raise QueryTimeoutError(str(exc)) from exc
            finally:
                timer.cancel()
        finally:
            conn.close()

        sources = tuple(
            self._sql_sources(tables[name], name, columns, rows) for name in guarded.tables
        )
        return QueryResult(
            columns=columns, rows=rows, sources=sources, truncated=len(rows) >= self._row_limit
        )

    @staticmethod
    def _sql_sources(
        entry: tuple[LoadedWorkbook, pl.DataFrame],
        name: str,
        columns: tuple[str, ...],
        rows: tuple[tuple[Any, ...], ...],
    ) -> SourceRange:
        wb, frame = entry
        sheet = wb.sheet_by_table[name]
        width = max(len(frame.columns) - 1, 1)
        last_col = get_column_letter(width)
        if ROW_COLUMN in columns and rows:
            idx = columns.index(ROW_COLUMN)
            numbers = [int(r[idx]) for r in rows if r[idx] is not None]
            if numbers:
                ref = f"A{min(numbers)}:{last_col}{max(numbers)}"
                return SourceRange(wb.file, sheet, ref, wb.document_id)
        first = int(cast(float, frame[ROW_COLUMN].min())) if frame.height else 2
        last = int(cast(float, frame[ROW_COLUMN].max())) if frame.height else 2
        return SourceRange(wb.file, sheet, f"A{first}:{last_col}{last}", wb.document_id)

    # ---- named ranges

    def read_named(self, workbook: LoadedWorkbook, name: str) -> tuple[Any, SourceRange]:
        for named in workbook.info.named_ranges:
            if named.name == name:
                if workbook.values is None:
                    raise UnknownNamedRangeError(name)
                value = workbook.values[named.sheet][named.ref].value
                return value, SourceRange(
                    workbook.file, named.sheet, named.ref, workbook.document_id
                )
        raise UnknownNamedRangeError(name)
