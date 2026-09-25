"""Workbook inspection with openpyxl (SPEC_04 §2, ADR-011): sheets, hidden state, header
columns, named ranges, macro presence, and how many formula cells lack a cached value.
Two loads per file — `data_only=True` for values, `data_only=False` to see formulas.
`keep_vba` is never set: VBA is neither loaded nor executed (SPEC_04 §1)."""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

MAX_HEADER_COLUMNS = 40
VBA_MEMBER = "xl/vbaProject.bin"


@dataclass(frozen=True)
class NamedRange:
    name: str
    sheet: str
    ref: str  # A1-style without `$`, e.g. "D14" or "B2:F21"


@dataclass(frozen=True)
class SheetInfo:
    name: str
    hidden: bool
    max_row: int
    max_col: int
    columns: tuple[str, ...]  # header row (row 1) values, as text


@dataclass(frozen=True)
class WorkbookInfo:
    file: str
    kind: str  # xlsx | xlsm | csv
    sheets: tuple[SheetInfo, ...]
    named_ranges: tuple[NamedRange, ...]
    has_macros: bool
    formula_cells: int
    formula_cells_without_cache: int
    notes: tuple[str, ...] = field(default_factory=tuple)

    @property
    def needs_recalculation(self) -> bool:
        return self.formula_cells_without_cache > 0


def has_vba(path: Path) -> bool:
    """`.xlsm` (or a renamed `.xlsx`) carrying a VBA project — detected from the zip
    members, never by trusting the extension."""
    try:
        with zipfile.ZipFile(path) as archive:
            return VBA_MEMBER in archive.namelist()
    except zipfile.BadZipFile:
        return False


def _header(ws: Any) -> tuple[str, ...]:
    values = []
    for col in range(1, min(ws.max_column, MAX_HEADER_COLUMNS) + 1):
        value = ws.cell(row=1, column=col).value
        values.append("" if value is None else str(value))
    return tuple(values)


def inspect_xlsx(path: Path) -> WorkbookInfo:
    formulas = load_workbook(path, data_only=False, read_only=False)
    values = load_workbook(path, data_only=True, read_only=False)
    sheets: list[SheetInfo] = []
    formula_cells = uncached = 0
    for ws in formulas.worksheets:
        ws_values = values[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells += 1
                    if ws_values[cell.coordinate].value is None:
                        uncached += 1
        sheets.append(
            SheetInfo(
                name=ws.title,
                hidden=ws.sheet_state != "visible",
                max_row=ws.max_row,
                max_col=ws.max_column,
                columns=_header(ws_values),
            )
        )
    named: list[NamedRange] = []
    for name, defined in formulas.defined_names.items():
        for sheet, ref in defined.destinations:
            named.append(NamedRange(name=name, sheet=sheet, ref=ref.replace("$", "")))
    kind = "xlsm" if has_vba(path) else "xlsx"
    return WorkbookInfo(
        file=path.name,
        kind=kind,
        sheets=tuple(sheets),
        named_ranges=tuple(sorted(named, key=lambda n: n.name)),
        has_macros=kind == "xlsm",
        formula_cells=formula_cells,
        formula_cells_without_cache=uncached,
    )


def read_csv_rows(path: Path, *, max_rows: int = 10_000) -> list[list[str]]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[list[str]] = []
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i >= max_rows:
            break
        rows.append(row)
    return rows


def inspect_csv(path: Path) -> WorkbookInfo:
    rows = read_csv_rows(path)
    header = tuple(rows[0]) if rows else ()
    width = max((len(r) for r in rows), default=0)
    sheet = SheetInfo(
        name=path.stem, hidden=False, max_row=len(rows), max_col=width, columns=header
    )
    return WorkbookInfo(
        file=path.name,
        kind="csv",
        sheets=(sheet,),
        named_ranges=(),
        has_macros=False,
        formula_cells=0,
        formula_cells_without_cache=0,
    )


def inspect_file(path: Path) -> WorkbookInfo:
    if path.suffix.lower() == ".csv":
        return inspect_csv(path)
    return inspect_xlsx(path)
