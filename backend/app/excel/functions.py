"""Predefined calculations (SPEC_04 §4): the LLM chooses the function and fills the
parameters, the code here reads the cells. Every result carries the cell/range it came
from. Named ranges are the contract with the generated workbooks
(`seed_data/generator/generate_excel.py`); a workbook without them simply cannot answer
through these functions (SQL remains)."""

from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date
from typing import Any, cast

import polars as pl
from openpyxl.utils import get_column_letter

from app.excel.calc import (
    ROW_COLUMN,
    CalcResult,
    CalculationEngine,
    ExcelError,
    LoadedWorkbook,
    SourceRange,
    UnknownNamedRangeError,
)
from app.excel.periods import normalize_period


class FunctionError(ExcelError):
    """Parameters the function cannot serve (unknown period, no workbook with the data)."""


@dataclass(frozen=True)
class FunctionSpec:
    name: str
    params: tuple[str, ...]
    description: str  # shown to the planning LLM, English
    run: Callable[[CalculationEngine, list[LoadedWorkbook], dict[str, Any]], CalcResult]


def _period(params: dict[str, Any], key: str = "period") -> str:
    raw = str(params.get(key) or "")
    period = normalize_period(raw)
    if period is None:
        raise FunctionError(f"unrecognised period {raw!r}")
    return period


def _first_named(
    engine: CalculationEngine, workbooks: list[LoadedWorkbook], name: str, *, prefer: str = ""
) -> tuple[Any, SourceRange, LoadedWorkbook]:
    ordered = sorted(workbooks, key=lambda wb: 0 if prefer and prefer in wb.file else 1)
    for wb in ordered:
        try:
            value, source = engine.read_named(wb, name)
        except UnknownNamedRangeError:
            continue
        return value, source, wb
    raise FunctionError(f"no workbook exposes {name}")


# S0-b: an ISO currency code written in a label, `"Closing (USD)"`, `"Variance (TRY)"`.
_CURRENCY_IN_LABEL = re.compile(r"\(([A-Z]{3})\)")


def currency_of(wb: LoadedWorkbook, source: SourceRange) -> str:
    """The currency of the cell `source` points at, read from the workbook itself — the
    column header (row 1) or, failing that, the nearest label to the left in the same row
    (the Covenant quarter sheets put the label in column B and the value in D). Never a
    constant in code (İş Emri 2026-10-13S0b: the schedule is USD, the code said EUR).
    Empty when no label carries a code — the figure is then served without a currency
    rather than with an invented one (rule 2)."""
    if wb.values is None:  # csv: no openpyxl workbook, no labels
        return ""
    ws = wb.values[source.sheet]
    cell = ws[source.ref.split(":")[0]]
    labels = [ws.cell(row=1, column=cell.column).value]
    labels += [ws.cell(row=cell.row, column=c).value for c in range(cell.column - 1, 0, -1)]
    for raw in labels:
        if isinstance(raw, str):
            match = _CURRENCY_IN_LABEL.search(raw)
            if match:
                return match.group(1)
    return ""


def _closing_column(frame: pl.DataFrame) -> str | None:
    """`closing_usd` / `closing_eur` / `closing` — whatever the Debt sheet's header says
    (S0-b: the column name was hard-coded as `closing_eur`, the regenerated workbook has
    `Closing (USD)`, and the function failed with "no debt schedule table")."""
    return next((c for c in frame.columns if c == "closing" or c.startswith("closing_")), None)


def _month_range(period: str) -> list[str]:
    """`Q2_2026` -> ["2026-04", "2026-05", "2026-06"]; `2026` -> 12 months; month -> itself."""
    if period.startswith("Q"):
        quarter, year = period.split("_")
        start = {"Q1": 1, "Q2": 4, "Q3": 7, "Q4": 10}[quarter]
        return [f"{year}-{m:02d}" for m in range(start, start + 3)]
    if len(period) == 4:
        return [f"{period}-{m:02d}" for m in range(1, 13)]
    return [period]


def _production_table(workbooks: list[LoadedWorkbook]) -> tuple[LoadedWorkbook, str, pl.DataFrame]:
    for wb in workbooks:
        for name, frame in wb.tables.items():
            if name.endswith("__production") and {"month", "mwh"} <= set(frame.columns):
                return wb, name, frame
    raise FunctionError("no production table")


def _production_rows(
    workbooks: list[LoadedWorkbook], period: str
) -> tuple[LoadedWorkbook, str, pl.DataFrame]:
    wb, name, frame = _production_table(workbooks)
    months = _month_range(period)
    subset = frame.filter(pl.col("month").cast(pl.Utf8).is_in(months))
    if subset.height == 0:
        raise FunctionError(f"no production rows for {period}")
    return wb, name, subset


def _as_int(value: Any) -> int:
    return int(cast(float, value))


def _range_of(wb: LoadedWorkbook, table: str, subset: pl.DataFrame, last_col: str) -> SourceRange:
    first = _as_int(subset[ROW_COLUMN].min())
    last = _as_int(subset[ROW_COLUMN].max())
    return SourceRange(
        wb.file, wb.sheet_by_table[table], f"A{first}:{last_col}{last}", wb.document_id
    )


def dscr(
    engine: CalculationEngine, workbooks: list[LoadedWorkbook], params: dict[str, Any]
) -> CalcResult:
    period = _period(params)
    if not period.startswith("Q"):
        raise FunctionError("dscr needs a quarter (e.g. Q2_2026)")
    value, source, _ = _first_named(engine, workbooks, f"DSCR_{period}", prefer="Covenant_Report")
    return CalcResult(value=float(value), unit="x", source=source, detail=f"DSCR {period}")


def outstanding_debt(
    engine: CalculationEngine, workbooks: list[LoadedWorkbook], params: dict[str, Any]
) -> CalcResult:
    raw = str(params.get("as_of") or "")
    period = normalize_period(raw) if raw else None
    if period is None or raw.lower() in {"today", "bugün", "demo_today", "güncel"}:
        # ADR-026: `Outstanding_DemoToday` is a value frozen at generation time, not a live
        # calculation. `Ledger_DemoToday` (same generation run) records which day it is for;
        # if the system's current day has moved on, the frozen figure is stale and must be
        # refused rather than served silently (rule 2/5).
        today = cast(date, params["_today"])
        baked_raw, _, _ = _first_named(
            engine, workbooks, "Ledger_DemoToday", prefer="Financial_Model"
        )
        baked_today = date.fromisoformat(str(baked_raw))
        if baked_today != today:
            raise FunctionError(
                f"Outstanding_DemoToday was computed for {baked_today.isoformat()}, not "
                f"today ({today.isoformat()}) — workbooks need regenerating"
            )
        value, source, wb = _first_named(
            engine, workbooks, "Outstanding_DemoToday", prefer="Financial_Model"
        )
        return CalcResult(
            value=float(value),
            unit=currency_of(wb, source),
            source=source,
            detail=f"outstanding on {today.isoformat()} (workbook snapshot)",
        )
    if period.startswith("Q"):
        value, source, wb = _first_named(
            engine, workbooks, f"Outstanding_{period}", prefer="Covenant_Report"
        )
        return CalcResult(
            value=float(value),
            unit=currency_of(wb, source),
            source=source,
            detail=f"outstanding after {period}",
        )
    # A date/month/year: last Debt-sheet row whose end date is on or before it.
    as_of = date.fromisoformat(f"{period}-01") if "-" in period else date(int(period), 12, 31)
    for wb in workbooks:
        for name, frame in wb.tables.items():
            closing = _closing_column(frame)
            if name.endswith("__debt") and closing and "end_date" in frame.columns:
                subset = frame.filter(pl.col("end_date").cast(pl.Date) <= as_of)
                if subset.height == 0:
                    raise FunctionError(f"no schedule row before {as_of}")
                row = subset.tail(1)
                excel_row = _as_int(row[ROW_COLUMN][0])
                # `_row` is column 0 of the frame, so a column's list index is its sheet
                # position (A = 1): the citation follows the sheet, not a fixed letter.
                letter = get_column_letter(frame.columns.index(closing))
                source = SourceRange(
                    wb.file, wb.sheet_by_table[name], f"{letter}{excel_row}", wb.document_id
                )
                return CalcResult(
                    float(row[closing][0]),
                    currency_of(wb, source),
                    source,
                    f"outstanding on {as_of}",
                )
    raise FunctionError("no debt schedule table")


def budget_variance(
    engine: CalculationEngine, workbooks: list[LoadedWorkbook], params: dict[str, Any]
) -> CalcResult:
    period = _period(params)
    if not period.startswith("Q"):
        raise FunctionError("budget_variance needs a quarter (e.g. Q2_2026)")
    value, source, wb = _first_named(
        engine, workbooks, f"Variance_{period}", prefer="Budget_vs_Actual"
    )
    budget, _, _ = _first_named(engine, workbooks, f"Budget_{period}", prefer="Budget_vs_Actual")
    actual, _, _ = _first_named(engine, workbooks, f"Actual_{period}", prefer="Budget_vs_Actual")
    detail = f"actual {float(actual):,.0f} - budget {float(budget):,.0f} ({period})"
    return CalcResult(float(value), currency_of(wb, source), source, detail)


def production(
    engine: CalculationEngine, workbooks: list[LoadedWorkbook], params: dict[str, Any]
) -> CalcResult:
    period = _period(params)
    if "-" in period:  # month — no named range for it means no data, never another month
        value, source, _ = _first_named(engine, workbooks, f"Production_{period.replace('-', '_')}")
        return CalcResult(float(value), "MWh", source, f"production {period}")
    if len(period) == 4:  # year
        try:
            value, source, _ = _first_named(engine, workbooks, f"MWh_Total_{period}")
            return CalcResult(float(value), "MWh", source, f"production {period}")
        except FunctionError:
            pass
    wb, table, subset = _production_rows(workbooks, period)
    total = float(subset["mwh"].cast(pl.Float64).sum())
    return CalcResult(
        total, "MWh", _range_of(wb, table, subset, "B"), f"sum of {subset.height} months ({period})"
    )


def capacity_factor(
    engine: CalculationEngine, workbooks: list[LoadedWorkbook], params: dict[str, Any]
) -> CalcResult:
    period = _period(params)
    wb, table, subset = _production_rows(workbooks, period)
    column = "capacity_factor" if "capacity_factor" in subset.columns else "capacity_factor_pct"
    if column not in subset.columns:
        column = next(c for c in subset.columns if c.startswith("capacity_factor"))
    value = float(cast(float, subset[column].cast(pl.Float64).mean()))
    detail = f"mean of {subset.height} month(s) ({period})"
    return CalcResult(round(value, 1), "%", _range_of(wb, table, subset, "D"), detail)


FUNCTIONS: dict[str, FunctionSpec] = {
    "dscr": FunctionSpec(
        "dscr", ("period",), "Debt service coverage ratio of a quarter (period like Q2_2026).", dscr
    ),
    "outstanding_debt": FunctionSpec(
        "outstanding_debt",
        ("as_of",),
        "Outstanding loan balance as of a quarter, month, year or 'today' (currency as in "
        "the workbook).",
        outstanding_debt,
    ),
    "budget_variance": FunctionSpec(
        "budget_variance",
        ("period",),
        "Actual minus budget for a quarter (currency as in the workbook).",
        budget_variance,
    ),
    "capacity_factor": FunctionSpec(
        "capacity_factor",
        ("period",),
        "Average capacity factor (%) of a month, quarter or year.",
        capacity_factor,
    ),
    "production": FunctionSpec(
        "production",
        ("period",),
        "Electricity production (MWh) of a month, quarter or year.",
        production,
    ),
}


def run_function(
    engine: CalculationEngine,
    workbooks: list[LoadedWorkbook],
    name: str,
    params: dict[str, Any],
    *,
    today: date,
) -> CalcResult:
    spec = FUNCTIONS.get(name)
    if spec is None:
        raise FunctionError(f"unknown function {name}")
    # `_today` is the system's own date (ADR-026), injected here — never part of the LLM's
    # plan JSON. Only `outstanding_debt` reads it (drift check against the baked snapshot).
    return spec.run(engine, workbooks, {**params, "_today": today})
