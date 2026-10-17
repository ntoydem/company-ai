"""İş Emri 2026-10-17P2B (ADR-037, plan §3.2b): header row detection, prelude and footer
note rows in `CachedValueEngine.load`. The Karatepe workbooks must load exactly as before
(compared against the pre-ADR-037 rule: header row 1, data from row 2); the repayment
schedules get their real header (Yeşilova row 4, Boztepe row 7) and typed columns."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

import polars as pl
import pytest
from openpyxl import Workbook, load_workbook

from app.excel.calc import (
    INFO_SUFFIX,
    CachedValueEngine,
    LoadedWorkbook,
    _frame,
    header_row,
    split_footer,
    table_name,
)
from app.excel.future_guard import future_row

EXCEL_DIR = Path(__file__).resolve().parent.parent / "seed_data" / "excel"
KARATEPE = (
    "Financial_Model_2026.xlsx",
    "Covenant_Report.xlsx",
    "Budget_vs_Actual_2026.xlsx",
    "Monthly_Production_2026.xlsx",
)
YESILOVA = "Yesilova_Repayment_Schedule_2026-09.xlsx"
YESILOVA_OLD = "Yesilova_Repayment_Schedule_2026-03.xlsx"
BOZTEPE = "Boztepe_Repayment_Schedule_84_Instalments.xlsx"
YSV_TABLE = "yesilova_repayment_schedule_2026_09__schedule"
BOZ_TABLE = "boztepe_repayment_schedule_84_instalments__schedule"
DATETIME = pl.Datetime("us")


@pytest.fixture(scope="module")
def engine() -> CachedValueEngine:
    return CachedValueEngine(timeout_s=10, row_limit=200)


@pytest.fixture(scope="module")
def loaded(engine: CachedValueEngine) -> dict[str, LoadedWorkbook]:
    return {f.name: engine.load(f) for f in sorted(EXCEL_DIR.glob("*.xlsx"))}


def _pre_adr037_tables(path: Path) -> dict[str, pl.DataFrame]:
    """The load rule until ADR-037, kept here as the reference: header = row 1."""
    tables: dict[str, pl.DataFrame] = {}
    for ws in load_workbook(path, data_only=True).worksheets:
        if ws.sheet_state != "visible" or ws.max_row < 2:
            continue
        rows = [
            (n, list(r))
            for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)
            if any(v is not None for v in r)
        ]
        if rows:
            tables[table_name(path.name, ws.title)] = _frame([c.value for c in ws[1]], rows)
    return tables


def _schema(frame: pl.DataFrame) -> list[tuple[str, pl.DataType]]:
    return list(frame.schema.items())


# ---------------------------------------------------------------- Karatepe unchanged


def test_karatepe_19_tables_are_identical_to_the_pre_adr037_load(
    loaded: dict[str, LoadedWorkbook],
) -> None:
    """Plan §3.2b / §8: every Karatepe sheet keeps row 1 as header — same table names,
    columns, types and cell values; no `_info` table appears."""
    count = 0
    for file in KARATEPE:
        before = _pre_adr037_tables(EXCEL_DIR / file)
        after = loaded[file].tables
        assert set(after) == set(before), file
        for name, frame in before.items():
            assert after[name].equals(frame), name
            assert _schema(after[name]) == _schema(frame), name
        count += len(before)
    assert count == 19


# ---------------------------------------------------------------- repayment schedules


def test_seed_workbook_column_snapshot(loaded: dict[str, LoadedWorkbook]) -> None:
    """Column snapshot of the three repayment schedule workbooks (Karatepe: test above)."""
    yesilova = [
        ("_row", pl.Int64),
        ("taksit_no", pl.Int64),
        ("vade", DATETIME),
        ("anapara_usd", pl.Int64),
        ("faiz_usd", pl.Int64),
        ("durum_bu_anl_k_g_r_nt_ye_g_re", pl.String),
    ]
    boztepe = [
        ("_row", pl.Int64),
        ("taksit_no", pl.Int64),
        ("vade", DATETIME),
        ("a_l_bakiyesi_usd", pl.Float64),
        ("anapara_usd", pl.Float64),
        ("faiz_usd", pl.Float64),
        ("toplam_deme_usd", pl.Float64),
        ("kapan_bakiyesi_usd", pl.Float64),
    ]
    info = [("_row", pl.Int64), ("key", pl.String), ("value", pl.String)]
    old_table = "yesilova_repayment_schedule_2026_03__schedule"
    expected = {
        YESILOVA: {YSV_TABLE: yesilova, YSV_TABLE + INFO_SUFFIX: info},
        YESILOVA_OLD: {old_table: yesilova, old_table + INFO_SUFFIX: info},
        BOZTEPE: {BOZ_TABLE: boztepe, BOZ_TABLE + INFO_SUFFIX: info},
    }
    for file, tables in expected.items():
        got = {name: _schema(frame) for name, frame in loaded[file].tables.items()}
        assert got == tables, file


def test_header_row_is_found_below_the_title_and_the_data_starts_under_it(
    loaded: dict[str, LoadedWorkbook],
) -> None:
    ysv = loaded[YESILOVA].tables[YSV_TABLE]
    boz = loaded[BOZTEPE].tables[BOZ_TABLE]
    assert ysv["_row"].min() == 5 and ysv.height == 24  # header row 4, instalments 1..24
    assert boz["_row"].min() == 8 and boz.height == 84  # header row 7, instalments 1..84
    assert ysv.filter(pl.col("taksit_no") == 12)["anapara_usd"].to_list() == [98000]
    assert ysv.filter(pl.col("taksit_no") == 11)["vade"].to_list() == [datetime(2026, 10, 14)]


def test_prelude_and_footer_notes_go_to_the_info_table_as_text(
    loaded: dict[str, LoadedWorkbook],
) -> None:
    """Nothing is lost: title, loan total, first due date (as a date, not a float) and the
    "Kontrol" / "Not" rows under the blank row are key/value text rows."""
    info = loaded[BOZTEPE].tables[BOZ_TABLE + INFO_SUFFIX]
    pairs = dict(zip(info["key"].to_list(), info["value"].to_list(), strict=True))
    assert info["_row"].to_list() == [1, 2, 3, 4, 5, 93, 94]
    assert pairs["Toplam Kredi (USD)"] == "15000000"
    assert pairs["Taksit Sayısı"] == "84"
    assert pairs["İlk Anapara Vade Tarihi"] == "2027-03-01"
    assert pairs["Kontrol: Son Kapanış Bakiyesi (0 olmalı)"] == "0"
    assert pairs["Not"] is not None and pairs["Not"].startswith("Son taksit (84)")
    ysv = loaded[YESILOVA].tables[YSV_TABLE + INFO_SUFFIX]
    assert ysv["_row"].to_list() == [1, 2, 30, 31]
    assert ysv["value"].to_list()[1] == "2026-09-01"
    assert loaded[BOZTEPE].sheet_by_table[BOZ_TABLE + INFO_SUFFIX] == "Schedule"


def test_sql_over_the_schedule_reads_typed_cells_and_cites_the_sheet(
    engine: CachedValueEngine, loaded: dict[str, LoadedWorkbook]
) -> None:
    """B-01 example 1 shape: one cell by instalment number — before ADR-037 the Boztepe
    column was text and `taksit_no = 12` raised a DuckDB conversion error."""
    ysv = engine.run_sql(
        [loaded[YESILOVA]], f"SELECT anapara_usd FROM {YSV_TABLE} WHERE taksit_no = 12"
    )
    assert ysv.rows == ((98000,),)
    boz = engine.run_sql(
        [loaded[BOZTEPE]], f"SELECT anapara_usd FROM {BOZ_TABLE} WHERE taksit_no = 12"
    )
    assert boz.rows == ((178571.43,),)
    assert boz.sources[0].label == f"{BOZTEPE} Schedule!A8:G91"


def test_future_guard_finds_the_due_date_column_of_a_schedule(
    loaded: dict[str, LoadedWorkbook],
) -> None:
    """`vade` is a date column now, so a row after TODAY is recognised (S2c guard)."""
    sql = f"SELECT anapara_usd FROM {BOZ_TABLE} ORDER BY vade DESC LIMIT 1"
    row = future_row([loaded[BOZTEPE]], sql, 178571.31, date(2026, 10, 6))
    assert row is not None and row.date_column == "vade" and row.excel_row == 91


# ---------------------------------------------------------------- rules (synthetic)


def test_header_rule_needs_three_text_cells_else_row_one() -> None:
    title: tuple[Any, ...] = ("Ödeme Planı", None, None)
    assert header_row([title, ("Taksit No", "Vade", "Anapara")]) == 2
    # Covenant quarter sheets: `Period | Q2_2026` (2 text cells) is a key/value row
    assert header_row([("Metric", "Value"), ("x", 1), ("Period", "Q2_2026")]) == 1
    # a row with a number among the cells is data, not a header
    assert header_row([title, ("Taksit No", "Vade", 3)]) == 1
    # blank strings do not count
    assert header_row([title, ("Taksit No", " ", "Anapara")]) == 1
    # only the first 10 rows are scanned
    assert header_row([title] * 10 + [("a", "b", "c")]) == 1
    assert header_row([]) == 1


def _save(tmp_path: Path, rows: list[list[Any]]) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for row in rows:
        ws.append(row)
    path = tmp_path / "sheet.xlsx"
    wb.save(path)
    return path


def test_a_second_data_block_after_a_blank_row_is_not_cut(
    tmp_path: Path, engine: CachedValueEngine
) -> None:
    """Naci 17.10 safety condition (b): a row after the blank row with ≥ half of the header
    columns as numbers/dates is data → nothing is cut ("kesilmedi, belirsiz")."""
    path = _save(
        tmp_path,
        [
            ["Plan"],
            ["No", "Vade", "Tutar", "Not"],
            [1, datetime(2027, 1, 1), 100, "a"],
            [None, None, None, None],
            [2, datetime(2027, 2, 1), None, None],  # 2 of 4 columns typed = half
            ["Not", "metin"],
        ],
    )
    loaded = engine.load(path)
    (main,) = [t for t in loaded.tables if not t.endswith(INFO_SUFFIX)]
    assert loaded.tables[main]["_row"].to_list() == [3, 5, 6]
    info = loaded.tables[main + INFO_SUFFIX]
    assert info["_row"].to_list() == [1]  # only the prelude


def test_split_footer_cuts_only_note_rows() -> None:
    data = (1, datetime(2027, 1, 1), 100.0, 5.0)
    blank = (None, None, None, None)
    notes = [(5, ("Not", "metin", None, None)), (6, ("Kontrol", 0, None, None))]
    body = [(2, data), (3, data), (4, blank), *notes]
    assert split_footer(body, 4) == ([(2, data), (3, data)], notes, False)
    # no blank row, or nothing after it → unchanged
    assert split_footer([(2, data), (3, data)], 4) == ([(2, data), (3, data)], [], False)
    assert split_footer([(2, data), (3, blank)], 4) == ([(2, data), (3, blank)], [], False)
    # a blank row before any data does not end the table
    assert split_footer([(2, blank), (3, data)], 4) == ([(2, blank), (3, data)], [], False)
    # a data-like row after the blank row → kept, flagged
    dated = (7, (None, datetime(2027, 3, 1), 3.0, None))
    assert split_footer([*body, dated], 4) == ([*body, dated], [], True)
    assert split_footer(body, 4)[2] is False


def test_footer_is_never_cut_on_a_sheet_whose_header_is_row_one(
    tmp_path: Path, engine: CachedValueEngine
) -> None:
    """Condition (a): Karatepe-style sheets (header row 1, blank rows between sections)
    keep every row in the table, as before."""
    path = _save(
        tmp_path,
        [["Metric", "Value", "Unit"], ["a", 1, "x"], [None, None, None], ["Not", "metin", None]],
    )
    loaded = engine.load(path)
    assert list(loaded.tables) == ["sheet__sheet"]
    assert loaded.tables["sheet__sheet"]["_row"].to_list() == [2, 4]
