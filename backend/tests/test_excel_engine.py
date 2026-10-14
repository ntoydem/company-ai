"""`CachedValueEngine` + predefined functions over the four committed demo workbooks
(`seed_data/excel/`, recalculated by LibreOffice, validated by `validate_excel.py`)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest

from app.excel.calc import (
    CachedValueEngine,
    LoadedWorkbook,
    NeedsRecalculationError,
    QueryTimeoutError,
    table_name,
)
from app.excel.functions import FunctionError, run_function
from app.excel.inspect import inspect_file
from app.excel.sql_guard import SqlRejectedError

EXCEL_DIR = Path(__file__).resolve().parent.parent / "seed_data" / "excel"
FILES = (
    "Financial_Model_2026.xlsx",
    "Covenant_Report.xlsx",
    "Budget_vs_Actual_2026.xlsx",
    "Monthly_Production_2026.xlsx",
)


@pytest.fixture(scope="module")
def engine() -> CachedValueEngine:
    return CachedValueEngine(timeout_s=10, row_limit=200)


@pytest.fixture(scope="module")
def workbooks(engine: CachedValueEngine) -> list[LoadedWorkbook]:
    return [engine.load(EXCEL_DIR / f) for f in FILES]


def test_all_committed_workbooks_load_without_error(engine: CachedValueEngine) -> None:
    """İş Emri 2026-10-12-F Adım 1: `FILES` (üstteki sabit) yalnız ilk 4 Karatepe
    workbook'unu kapsıyordu — Yeşilova/Boztepe eklendiğinde bu teste hiç dahil
    edilmemişlerdi, bu yüzden `_frame()`'in Yeşilova'nın "Schedule" sayfasındaki
    düzensiz önsöz satırıyla (başlık + anlık görüntü tarihi, gerçek tablo başlığından
    önce) çöktüğü (`polars.exceptions.ComputeError`) hiç yakalanmamıştı. Artık
    `seed_data/excel/`'deki HER workbook burada yükleniyor."""
    for path in sorted(EXCEL_DIR.glob("*.xlsx")):
        engine.load(path)  # exception fırlatmamalı


def test_generated_workbooks_have_no_uncached_formula_and_a_hidden_meta_sheet() -> None:
    """PHASES.md 4.2 kabul kriteri 1: cached değerler dolu (`data_only` boş hücre yok)."""
    for file in FILES:
        info = inspect_file(EXCEL_DIR / file)
        assert info.formula_cells > 0, file
        assert info.formula_cells_without_cache == 0, file
        assert any(s.name == "_meta" and s.hidden for s in info.sheets), file
        assert not info.has_macros


def test_dscr_q2_2026_comes_from_the_covenant_report_cell(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    """PHASES.md 4.2 kabul kriteri 2: "Ankara RES 2026 Q2 DSCR kaç?" ->
    Covenant_Report.xlsx Q2_2026!D14."""
    result = run_function(engine, workbooks, "dscr", {"period": "2026 Q2"}, today=date(2026, 10, 6))
    assert result.value == pytest.approx(1.37)
    assert result.source.label == "Covenant_Report.xlsx Q2_2026!D14"


def test_outstanding_debt_today_and_after_a_quarter(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    today = run_function(
        engine, workbooks, "outstanding_debt", {"as_of": "today"}, today=date(2026, 10, 6)
    )
    assert today.value == pytest.approx(11_671_800)
    assert today.source.file == "Financial_Model_2026.xlsx" and today.source.sheet == "Debt"
    q = run_function(
        engine, workbooks, "outstanding_debt", {"as_of": "Q4_2024"}, today=date(2026, 10, 6)
    )
    assert q.value == pytest.approx(12_703_700)
    assert q.source.label == "Covenant_Report.xlsx Q4_2024!D12"


def test_outstanding_debt_today_refuses_a_stale_snapshot(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    """ADR-026: `Outstanding_DemoToday` was frozen for 06.10.2026 (`Ledger_DemoToday`); once
    the system's own day has moved on, the figure must be refused, not served silently."""
    with pytest.raises(FunctionError):
        run_function(
            engine, workbooks, "outstanding_debt", {"as_of": "today"}, today=date(2026, 10, 7)
        )


def test_outstanding_debt_unit_comes_from_the_workbook_label_not_from_code(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    """İş Emri 2026-10-13S0b A/2 (BT-01 E5): the schedule is in USD (`Closing (USD)`,
    `Outstanding after half (USD)`); the code used to say EUR for every branch."""
    today = run_function(
        engine, workbooks, "outstanding_debt", {"as_of": "today"}, today=date(2026, 10, 6)
    )
    assert today.value == pytest.approx(11_671_800) and today.unit == "USD"
    quarter = run_function(
        engine, workbooks, "outstanding_debt", {"as_of": "Q4_2024"}, today=date(2026, 10, 6)
    )
    assert quarter.unit == "USD" and quarter.source.label == "Covenant_Report.xlsx Q4_2024!D12"
    variance = run_function(
        engine, workbooks, "budget_variance", {"period": "Q2 2026"}, today=date(2026, 10, 6)
    )
    assert variance.unit == "TRY"  # `Variance (TRY)` header, no constant either


def test_outstanding_debt_on_a_date_resolves_the_closing_column_from_the_sheet(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    """İş Emri 2026-10-13S0b A/3: a dated question looked for `closing_eur`; the Debt
    sheet has `Closing (USD)` → `closing_usd`, so every dated call failed with "no debt
    schedule table". The column (and its letter, H) now come from the sheet."""
    september = run_function(
        engine, workbooks, "outstanding_debt", {"as_of": "2026-09"}, today=date(2026, 10, 6)
    )
    assert september.value == pytest.approx(11_671_800) and september.unit == "USD"
    assert september.source.label == "Financial_Model_2026.xlsx Debt!H9"  # H1_2026 row
    assert september.detail == "outstanding on 2026-09-01"
    tr_date = run_function(
        engine, workbooks, "outstanding_debt", {"as_of": "15.09.2026"}, today=date(2026, 10, 6)
    )
    assert tr_date.value == pytest.approx(11_671_800)
    year = run_function(
        engine, workbooks, "outstanding_debt", {"as_of": "2024"}, today=date(2026, 10, 6)
    )
    assert year.value == pytest.approx(12_703_700)  # = Covenant Q4_2024 outstanding
    assert year.source.label == "Financial_Model_2026.xlsx Debt!H6"
    with pytest.raises(FunctionError):  # before the first drawdown: no row, no guess
        run_function(
            engine, workbooks, "outstanding_debt", {"as_of": "2022-06"}, today=date(2026, 10, 6)
        )


def test_production_of_a_month_without_data_is_refused_not_another_month(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    """İş Emri 2026-10-13S0b A/1 (BT-01 E1): "last month" from 06.10.2026 is 2026-09; the
    workbook ends at 2026-08 (13.538 MWh). The engine says "no data" — it never serves
    June (10.951) or any other row instead."""
    with pytest.raises(FunctionError):
        run_function(
            engine, workbooks, "production", {"period": "2026-09"}, today=date(2026, 10, 6)
        )
    latest = run_function(
        engine, workbooks, "production", {"period": "2026-08"}, today=date(2026, 10, 6)
    )
    assert latest.value == pytest.approx(13_538)


def test_budget_variance_production_and_capacity_factor(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    variance = run_function(
        engine, workbooks, "budget_variance", {"period": "Q2 2026"}, today=date(2026, 10, 6)
    )
    assert variance.value == pytest.approx(882_000)
    assert variance.unit == "TRY" and variance.source.sheet == "Summary"

    month = run_function(
        engine, workbooks, "production", {"period": "2026-08"}, today=date(2026, 10, 6)
    )
    assert month.value == pytest.approx(13_538)
    year = run_function(
        engine, workbooks, "production", {"period": "2025"}, today=date(2026, 10, 6)
    )
    assert year.value == pytest.approx(164_244)
    assert year.source.label == "Monthly_Production_2026.xlsx KPI!B4"
    quarter = run_function(
        engine, workbooks, "production", {"period": "Q2_2026"}, today=date(2026, 10, 6)
    )
    assert quarter.value == pytest.approx(12_561 + 12_304 + 10_951)
    assert quarter.source.sheet == "Production" and quarter.source.ref.startswith("A")

    cf = run_function(
        engine, workbooks, "capacity_factor", {"period": "2026-01"}, today=date(2026, 10, 6)
    )
    assert cf.value == pytest.approx(38.2)


def test_function_errors_on_bad_period_or_unknown_name(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    with pytest.raises(FunctionError):
        run_function(engine, workbooks, "dscr", {"period": "some day"}, today=date(2026, 10, 6))
    with pytest.raises(FunctionError):
        run_function(engine, workbooks, "nope", {}, today=date(2026, 10, 6))


def test_sql_sum_over_the_production_table_cites_the_sheet_range(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    table = table_name("Monthly_Production_2026.xlsx", "Production")
    result = engine.run_sql(
        workbooks, f"SELECT SUM(mwh) AS total FROM {table} WHERE month LIKE '2026-%'"
    )
    assert result.columns == ("total",)
    assert result.rows[0][0] == pytest.approx(108_858)
    assert len(result.sources) == 1
    assert result.sources[0].file == "Monthly_Production_2026.xlsx"
    assert result.sources[0].sheet == "Production"
    assert result.sources[0].ref.startswith("A2:")


def test_sql_row_column_narrows_the_cited_range(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    table = table_name("Covenant_Report.xlsx", "Summary")
    result = engine.run_sql(
        workbooks, f"SELECT _row, quarter, dscr FROM {table} WHERE quarter = 'Q2_2026'"
    )
    assert result.rows[0][1:] == ("Q2_2026", pytest.approx(1.37))
    assert result.sources[0].ref == "A12:E12"


def test_sql_guard_and_external_access_are_both_enforced(
    workbooks: list[LoadedWorkbook], engine: CachedValueEngine
) -> None:
    with pytest.raises(SqlRejectedError):
        engine.run_sql(workbooks, "SELECT * FROM read_csv('/etc/passwd')")
    with pytest.raises(SqlRejectedError):
        engine.run_sql(workbooks, "DROP TABLE covenant_report__summary")


def test_sql_timeout_interrupts_a_runaway_query(workbooks: list[LoadedWorkbook]) -> None:
    fast = CachedValueEngine(timeout_s=0.2, row_limit=10)
    table = table_name("Monthly_Production_2026.xlsx", "Production")
    with pytest.raises(QueryTimeoutError):
        fast.run_sql(
            workbooks,
            f"SELECT COUNT(*) FROM {table} a, {table} b, {table} c, {table} d, "
            f"{table} e, {table} f",
        )


def test_workbook_without_cached_values_is_refused(
    tmp_path: Path, engine: CachedValueEngine
) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = "x", 1, "=A2*2"
    path = tmp_path / "raw.xlsx"
    wb.save(path)
    with pytest.raises(NeedsRecalculationError):
        engine.load(path)


def test_a_column_starting_with_dates_then_text_does_not_crash(
    tmp_path: Path, engine: CachedValueEngine
) -> None:
    """İş Emri 2026-10-12-F Adım 1 kök nedeni: Yeşilova'nın "Schedule" sayfası, gerçek
    tablo başlığından (`ws[4]`) önce bir başlık + anlık görüntü tarihi satırı taşıyordu;
    `CachedValueEngine.load()` her zaman `ws[1]`'i başlık sayıp 2. satırdan veri okuduğu
    için bu önsöz, A sütununda `datetime` (satır 2: anlık görüntü tarihi) → `str`
    (satır 4: "Taksit No" metni, veri gibi okunmuş) sırasıyla karışık bir tip üretiyor —
    polars `strict=False`'a rağmen `ComputeError('expected duration or datetime, got
    str')` fırlatıyordu. `/api/ask`'ta `finans`/`yonetim` kullanıcısının HER veri sorusu
    erişebildiği TÜM workbook'ları yüklediği için bu, LLM sağlayıcısından bağımsız,
    deterministik bir HTTP 500'dü (bkz. rapor)."""
    from datetime import datetime

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Başlık"
    ws["A2"] = datetime(2026, 3, 15)
    ws["A3"] = "Taksit No"
    ws["A4"] = 1
    ws["B1"], ws["B2"], ws["B3"], ws["B4"] = "x", "x", "x", "x"
    path = tmp_path / "mixed.xlsx"
    wb.save(path)
    loaded = engine.load(path)  # exception fırlatmamalı
    (table_name_,) = loaded.tables
    result = engine.run_sql([loaded], f"SELECT COUNT(*) FROM {table_name_}")
    assert result.rows[0][0] == 3  # satır 2, 3, 4 (başlık sayılan satır 1 hariç)


def test_csv_is_loaded_as_a_single_table(tmp_path: Path, engine: CachedValueEngine) -> None:
    path = tmp_path / "opex.csv"
    path.write_text("period;amount\nQ1_2026;100\nQ2_2026;250\n", encoding="utf-8")
    loaded = engine.load(path)
    assert list(loaded.tables) == ["opex__opex"]
    result = engine.run_sql([loaded], "SELECT SUM(CAST(amount AS DOUBLE)) FROM opex__opex")
    assert result.rows[0][0] == pytest.approx(350)
    assert loaded.info.kind == "csv" and not loaded.info.has_macros


# ---------------------------------------------------------------- S2c: a row after TODAY


DEBT = "financial_model_2026__debt"
DEBT_LATEST = f"SELECT closing_usd FROM {DEBT} ORDER BY end_date DESC LIMIT 1"


def test_future_row_identifies_a_debt_row_dated_after_today(
    workbooks: list[LoadedWorkbook],
) -> None:
    """İş Emri 2026-10-14S2c (B2-b §4c): `ORDER BY end_date DESC LIMIT 1` on the Debt sheet
    reads Closing = 0 from H2_2034 — the guard names that row (table, date column, Excel
    row, label), without any figure; the past row H1_2026 (11.671.800) is not flagged."""
    from app.excel.future_guard import future_row

    today = date(2026, 10, 6)
    row = future_row(workbooks, DEBT_LATEST, 0, today)
    assert row is not None
    assert row.table == DEBT and row.date_column == "end_date"
    assert row.excel_row == 26 and row.label == "H2_2034"
    # the right row (H1_2026, end date 30.06.2026) — no guard
    assert future_row(workbooks, DEBT_LATEST, 11_671_800, today) is None
    # explicit future period: the engine still reports it; the plan's `future_ok` decides
    by_period = f"SELECT closing_usd FROM {DEBT} WHERE period = 'H2_2030'"
    assert future_row(workbooks, by_period, 6_725_900, today) is not None


def test_future_row_stays_silent_when_it_cannot_attribute_the_value(
    workbooks: list[LoadedWorkbook],
) -> None:
    """No date column (Production: `month` is text), a value in no row, a column the SQL
    does not name, an aggregate, or a value that also occurs in a past row → `None`."""
    from app.excel.future_guard import future_row

    today = date(2026, 10, 6)
    production = "SELECT mwh FROM monthly_production_2026__production ORDER BY month DESC LIMIT 1"
    assert future_row(workbooks, production, 13_538, today) is None
    assert future_row(workbooks, DEBT_LATEST, 12_345, today) is None
    assert future_row(workbooks, f"SELECT SUM(interest_usd) FROM {DEBT}", 0, today) is None
    # `drawdown_usd` is 0 in past rows too — cannot be attributed to a future row
    drawdown = f"SELECT drawdown_usd FROM {DEBT} ORDER BY end_date DESC LIMIT 1"
    assert future_row(workbooks, drawdown, 0, today) is None
    assert future_row(workbooks, "not a select", 0, today) is None
    dated = f"SELECT end_date FROM {DEBT} ORDER BY end_date DESC LIMIT 1"
    assert future_row(workbooks, dated, date(2034, 12, 31), today) is None
