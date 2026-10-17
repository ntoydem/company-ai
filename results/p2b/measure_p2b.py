"""P2B measure (eval DB, read-only, LLM-free): (1) table columns of every Excel document,
(2) repayment-schedule cells before/after ADR-037, (3) the P2A acceptance measure again —
never another project's figure.

Run (eval project only; `before` with the pre-P2B tree mounted at /app):
  docker compose -p company-ai-eval --project-directory . -f infra/docker-compose.yml
  --env-file infra/.env.eval run --rm -T --no-deps [-v <tree>/backend:/app] backend
  python - before|after < results/p2b/measure_p2b.py"""

import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select, text

from app.api.deps import get_calculation_engine
from app.core.config import get_settings
from app.core.db import get_session_factory
from app.excel import calc
from app.models.document import Document
from app.repositories import user_repo
from app.schemas.excel import ExcelAskRequest
from app.services.excel_ask import PLAN_SYSTEM_PROMPT, answer_data_question
from app.services.llm import LLMRequest, LLMResponse

LABEL = sys.argv[1] if len(sys.argv) > 1 else "run"
AFTER = hasattr(calc, "header_row")
YSV = "yesilova_repayment_schedule_2026_09__schedule"
BOZ = "boztepe_repayment_schedule_84_instalments__schedule"
OLD_YSV_KEY = "deme_plan_ye_ilova_res"
OLD_BOZ_KEY = "boztepe_res_anapara_geri_deme_plan_84_taksit"


def sql(statement: str) -> dict[str, Any]:
    return {"kind": "sql", "sql": statement, "future_ok": True}


def ysv(column_old: str, column_new: str, instalment: int) -> dict[str, Any]:
    if AFTER:
        return sql(f"SELECT {column_new} FROM {YSV} WHERE taksit_no = {instalment}")
    return sql(f"SELECT {column_old} FROM {YSV} WHERE {OLD_YSV_KEY} = '{instalment}'")


def boz(column_old: str, column_new: str, instalment: int) -> dict[str, Any]:
    if AFTER:
        return sql(f"SELECT {column_new} FROM {BOZ} WHERE taksit_no = {instalment}")
    return sql(f"SELECT {column_old} FROM {BOZ} WHERE {OLD_BOZ_KEY} = '{instalment}'")


NEXT_YSV = (
    sql(f"SELECT vade FROM {YSV} WHERE vade > DATE '2026-10-06' ORDER BY vade LIMIT 1")
    if AFTER
    else sql(f"SELECT col_b FROM {YSV} WHERE col_b > '2026-10-06' ORDER BY col_b LIMIT 1")
)
# (id, question, plan) — part 2, the B-01 cells
CELLS = [
    (
        "Y12A",
        "Yeşilova RES ödeme planında 12. taksitin anaparası ne kadar?",
        ysv("col_c", "anapara_usd", 12),
    ),
    ("Y11V", "Yeşilova RES ödeme planında 11. taksitin vadesi ne?", ysv("col_b", "vade", 11)),
    (
        "Y11F",
        "Yeşilova RES ödeme planında 11. taksitin faizi ne kadar?",
        ysv("col_d", "faiz_usd", 11),
    ),
    ("YNXT", "Yeşilova RES'in bir sonraki taksitinin vadesi ne?", NEXT_YSV),
    (
        "B12A",
        "Boztepe RES ödeme planında 12. taksitin anaparası ne kadar?",
        boz("col_d", "anapara_usd", 12),
    ),
    ("B12V", "Boztepe RES ödeme planında 12. taksitin vadesi ne?", boz("col_b", "vade", 12)),
    (
        "B12F",
        "Boztepe RES ödeme planında 12. taksitin faizi ne kadar?",
        boz("col_e", "faiz_usd", 12),
    ),
]
TODAY_DEBT = {"kind": "function", "name": "outstanding_debt", "params": {"as_of": "today"}}
DSCR = {"kind": "function", "name": "dscr", "params": {"period": "Q2_2026"}}
KARATEPE_DEBT_SQL = {
    "kind": "sql",
    "sql": "SELECT closing_usd FROM financial_model_2026__debt "
    "WHERE end_date <= DATE '2026-10-06' ORDER BY end_date DESC LIMIT 1",
}
YSV_12 = ysv("col_c", "anapara_usd", 12)
BOZ_12 = boz("col_d", "anapara_usd", 12)
# part 3 — the P2A cases (results/p2a/measure_p2a.py); Y4/B4/U2 use this tree's columns
CASES = [
    ("Y1", "Yeşilova RES'in kalan kredi borcu ne kadar?", "YSV_RES", TODAY_DEBT),
    ("Y2", "Yeşilova RES'in güncel DSCR değeri kaç?", "YSV_RES", DSCR),
    ("Y3", "Yeşilova RES'in kalan kredi borcu ne kadar?", "YSV_RES", KARATEPE_DEBT_SQL),
    ("Y4", "Yeşilova RES ödeme planında 12. taksitin anaparası ne kadar?", "YSV_RES", YSV_12),
    ("B1", "Boztepe RES'in kalan kredi borcu ne kadar?", "BOZ_RES", TODAY_DEBT),
    ("B2", "Boztepe RES'in güncel DSCR değeri kaç?", "BOZ_RES", DSCR),
    ("B3", "Boztepe RES'in kalan kredi borcu ne kadar?", "BOZ_RES", KARATEPE_DEBT_SQL),
    ("B4", "Boztepe RES ödeme planında 12. taksitin anaparası ne kadar?", "BOZ_RES", BOZ_12),
    ("K1", "Karatepe RES'in kalan kredi borcu ne kadar?", "ANK_RES", TODAY_DEBT),
    ("K2", "Karatepe RES'in güncel DSCR değeri kaç?", "ANK_RES", DSCR),
    ("U1", "Kalan kredi borcu ne kadar?", None, TODAY_DEBT),
    ("U2", "Ödeme planında 12. taksitin anaparası ne kadar?", None, YSV_12),
]
USERS = ("finans", "yonetim")


class PlanOnlyLLM:
    """Planner → the fixed plan; answering model → "-", so the template answer shows."""

    def __init__(self, plan: dict[str, Any]) -> None:
        self.plan = plan
        self.calls: list[str] = []

    def complete(self, request: LLMRequest) -> LLMResponse:
        self.calls.append("plan" if request.system == PLAN_SYSTEM_PROMPT else "answer")
        reply = json.dumps(self.plan) if request.system == PLAN_SYSTEM_PROMPT else "-"
        return LLMResponse(
            text=reply,
            model="fake",
            tokens_in=0,
            tokens_out=0,
            tokens_reasoning=None,
            finish_reason="stop",
            latency_ms=0,
        )


def sheet_cut_status(path: Path) -> dict[str, str]:
    """ADR-037 footer rule per visible sheet (after only): header row and cut status."""
    out: dict[str, str] = {}
    for ws in load_workbook(path, data_only=True).worksheets:
        if ws.sheet_state != "visible" or ws.max_row < 2:
            continue
        top = list(ws.iter_rows(max_row=calc.HEADER_SCAN_ROWS, values_only=True))
        first = calc.header_row(top)
        status = "başlık 1. satır (kesme kuralı uygulanmaz)"
        if first > 1:
            body = list(enumerate(ws.iter_rows(min_row=first + 1, values_only=True), first + 1))
            width = sum(calc._filled(v) for v in top[first - 1])
            _, notes, kept = calc.split_footer(body, width)
            status = (
                "kesilmedi, belirsiz"
                if kept
                else (f"kesildi ({len(notes)} not satırı)" if notes else "alt not yok")
            )
        out[ws.title] = f"header row {first}; {status}"
    return out


def run(session: Any, user: Any, question: str, plan: dict[str, Any], settings: Any) -> Any:
    llm = PlanOnlyLLM(plan)
    result = answer_data_question(
        session,
        user,
        ExcelAskRequest(question=question),
        llm,  # type: ignore[arg-type]
        settings,
        engine,
        write_audit=False,
        original_question=question,
    )
    return result, llm.calls


settings = get_settings()
engine = get_calculation_engine()
report: dict[str, Any] = {"label": LABEL, "after_code": AFTER}
with get_session_factory()() as session:
    session.execute(text("SET SESSION CHARACTERISTICS AS TRANSACTION READ ONLY"))
    session.commit()
    print(
        "db",
        session.execute(text("select current_database()")).scalar(),
        "read_only",
        session.execute(text("show transaction_read_only")).scalar(),
    )
    # 1. columns
    columns: dict[str, Any] = {}
    documents = session.scalars(select(Document).order_by(Document.file_name)).all()
    for document in documents:
        if not (document.file_name or "").endswith(".xlsx"):
            continue
        path = settings.documents_dir / document.storage_path
        wb = engine.load(path, document_id=document.id, file_name=document.file_name)
        columns[document.file_name] = {
            "project": document.project.code if document.project else "GROUP",
            "status": str(document.status),
            "tables": {
                name: [f"{c}:{t}" for c, t in frame.schema.items()]
                for name, frame in wb.tables.items()
            },
            "sheets": sheet_cut_status(path) if AFTER else {},
        }
    report["columns"] = columns
    finans = user_repo.get_by_username(session, "finans")
    assert finans is not None
    # 2. B-01 cells, with and without the Ek-F formatter
    cells = []
    for case_id, question, plan in CELLS:
        for ek_f in (False, True):
            result, calls = run(
                session, finans, question, plan, settings.model_copy(update={"ek_f_enabled": ek_f})
            )
            cells.append(
                {
                    "id": case_id,
                    "ek_f": ek_f,
                    "sql": plan["sql"],
                    "answered": result.answered,
                    "value": repr(result.value),
                    "value_type": type(result.value).__name__,
                    "formatted": result.formatted_value,
                    "answer": result.answer,
                    "sources": [s.label for s in result.sources],
                    "llm_calls": calls,
                }
            )
            print(json.dumps(cells[-1], ensure_ascii=False))
    report["cells"] = cells
    # 3. P2A acceptance again
    rows = []
    for username in USERS:
        user = user_repo.get_by_username(session, username)
        assert user is not None, username
        for case_id, question, expected, plan in CASES:
            result, calls = run(session, user, question, plan, settings)
            projects = []
            for card in result.sources:
                document = session.get(Document, card.document_id)
                projects.append(document.project.code if document and document.project else "GROUP")
            foreign = bool(expected) and result.answered and any(p != expected for p in projects)
            rows.append(
                {
                    "user": username,
                    "id": case_id,
                    "answered": result.answered,
                    "value": result.formatted_value,
                    "source_projects": projects,
                    "project_question": result.project_question,
                    "answer": result.answer,
                    "llm_calls": calls,
                    "foreign_project_value": foreign,
                }
            )
            print(json.dumps(rows[-1], ensure_ascii=False))
    report["p2a"] = rows
    report["foreign_total"] = sum(r["foreign_project_value"] for r in rows)
    session.rollback()
print("FOREIGN_TOTAL", report["foreign_total"], "of", len(report["p2a"]))
with open(f"results/p2b/measure_{LABEL}.json", "w") as fh:
    json.dump(report, fh, ensure_ascii=False, indent=1)
